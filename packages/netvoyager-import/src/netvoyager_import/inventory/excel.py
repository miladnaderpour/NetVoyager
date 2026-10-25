
"""Import normalized device observations from an Excel inventory.

Excel inventory rows are parsed and validated independently of canonical
network device identity.

The importer provides two public entry points:

- read_excel_inventory(): detailed parsed inventory and validation issues.
- read_network_devices(): normalized SourceDataset for device discovery.

The importer does not create canonical devices, resolve identities, or make
authorization decisions. Those responsibilities belong to Analysis.
"""

from __future__ import annotations

import logging
from collections import Counter
from collections.abc import Collection, Mapping
from dataclasses import dataclass, replace
from ipaddress import IPv4Address
from pathlib import Path
from types import MappingProxyType
from typing import Any, Literal, TypeAlias
from uuid import NAMESPACE_URL, uuid5

from openpyxl import load_workbook

from netvoyager_analysis.devices.observations import DeviceObservation
from netvoyager_analysis.sources.models import Source, SourceDataset


logger = logging.getLogger("netvoyager.import.inventory.excel")


# ---------------------------------------------------------------------------
# Source identity
# ---------------------------------------------------------------------------

EXCEL_SOURCE_NAME = "excel_inventory"
EXCEL_SOURCE_TYPE = "device_inventory"

EXCEL_SOURCE_ID = uuid5(
    NAMESPACE_URL,
    "netvoyager:source:excel_inventory",
)


# ---------------------------------------------------------------------------
# Excel column definitions
# ---------------------------------------------------------------------------

NAME_COLUMN = "name"
HOST_NAME_COLUMN = "host_name"
SITE_COLUMN = "site"
ROLE_COLUMN = "role"

PRIMARY_IPV4_COLUMN = "primary_ipv4"
MANAGEMENT_IP_COLUMN = "management_ip"

MANUFACTURER_COLUMN = "manufacturer"
MODEL_COLUMN = "model"
SERIAL_NUMBER_COLUMN = "serial_number"
ASSET_TAG_COLUMN = "asset_tag"
SOFTWARE_VERSION_COLUMN = "software_version"

STATUS_COLUMN = "status"
VALIDATION_STATUS_COLUMN = "validation_status"


DuplicateNamePolicy: TypeAlias = Literal[
    "prefix",
    "warn",
    "error",
    "ignore",
]

DuplicateSerialPolicy: TypeAlias = Literal[
    "warn",
    "error",
    "ignore",
]


UNKNOWN_SERIAL_VALUES = frozenset({
    "n/a",
    "na",
    "none",
    "null",
    "unknown",
    "0",
    "-",
})


# ---------------------------------------------------------------------------
# Excel parsing models
# ---------------------------------------------------------------------------

@dataclass(frozen=True, slots=True)
class ExcelDeviceObservation:
    """Device attributes reported by an individual Excel inventory row.

    This is an intermediate import representation, not a canonical device.

    Original source values are retained for validation and diagnostics.
    """

    name: str | None = None
    host_name: str | None = None
    site: str | None = None
    role: str | None = None

    primary_IPv4: IPv4Address | None = None
    management_ip: IPv4Address | None = None

    manufacturer: str | None = None
    model: str | None = None
    serial_number: str | None = None
    asset_tag: str | None = None
    software_version: str | None = None

    status: str | None = None
    validation_status: str | None = None

    line_number: int | None = None
    raw_line: str | None = None


@dataclass(frozen=True, slots=True)
class InventoryIssue:
    """A validation or consistency issue in the Excel inventory."""

    row_number: int
    reason: str


@dataclass(frozen=True, slots=True)
class InventoryReadResult:
    """Detailed result of Excel inventory parsing.

    Observations preserve source-row order and are never merged.

    Management IP mappings retain every site reporting an address,
    including conflicting and duplicated addresses.
    """

    observations: tuple[ExcelDeviceObservation, ...]
    issues: tuple[InventoryIssue, ...]
    ip_sites_mapping: Mapping[IPv4Address, frozenset[str]]


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

def _text(value: Any) -> str | None:
    """Normalize an Excel cell to stripped text or None."""
    if value is None:
        return None

    result = str(value).strip()
    return result or None


def _cell(
    row: tuple[Any, ...],
    columns: Mapping[str, int],
    column_name: str,
) -> Any:
    """Read a named column, returning None when unavailable."""
    index = columns.get(column_name)

    if index is None or index >= len(row):
        return None

    return row[index]


def _ipv4(
    value: Any,
    *,
    row_number: int,
    field_name: str,
    issues: list[InventoryIssue],
    required: bool = False,
) -> IPv4Address | None:
    """Parse an IPv4 address without discarding an invalid source row."""
    text = _text(value)

    if text is None:
        if required:
            issues.append(
                InventoryIssue(
                    row_number=row_number,
                    reason=f"Missing {field_name}",
                )
            )
        return None

    try:
        return IPv4Address(text)
    except ValueError:
        issues.append(
            InventoryIssue(
                row_number=row_number,
                reason=f"Invalid {field_name}",
            )
        )
        return None


def _raw_line(
    row: tuple[Any, ...],
    header: tuple[Any, ...],
) -> str:
    """Serialize populated source cells for diagnostics and provenance."""
    values: list[str] = []

    for column_name, value in zip(header, row):
        name = _text(column_name)
        text = _text(value)

        if name is not None and text is not None:
            values.append(f"{name}={text}")

    return " | ".join(values)


def _duplicate_key(value: str | None) -> str | None:
    """Normalize a device identifier for case-insensitive comparison."""
    if value is None:
        return None

    return value.strip().casefold() or None


def _serial_key(value: str | None) -> str | None:
    """Normalize a serial number while excluding placeholder values."""
    key = _duplicate_key(value)

    if key is None or key in UNKNOWN_SERIAL_VALUES:
        return None

    return key


def _normalize_filter(
    values: Collection[str] | None,
) -> frozenset[str] | None:
    """Normalize an optional site/device selection filter."""
    if values is None:
        return None

    return frozenset(
        normalized
        for value in values
        if (normalized := _duplicate_key(value)) is not None
    )


# ---------------------------------------------------------------------------
# Duplicate handling
# ---------------------------------------------------------------------------

def _resolve_duplicate_observations(
    observations: list[ExcelDeviceObservation],
    issues: list[InventoryIssue],
    *,
    duplicate_name_policy: DuplicateNamePolicy,
    duplicate_serial_policy: DuplicateSerialPolicy,
) -> list[ExcelDeviceObservation]:
    """Apply duplicate-name and duplicate-serial policies.

    Duplicate names are checked across the complete workbook.

    Under the prefix policy, duplicated names receive deterministic
    DUP-row-name identifiers.

    Duplicate serials produce warnings or errors according to policy.
    No records are merged or silently discarded.
    """
    name_counts = Counter(
        key
        for observation in observations
        if (key := _duplicate_key(observation.name)) is not None
    )

    serial_counts = Counter(
        key
        for observation in observations
        if (key := _serial_key(observation.serial_number)) is not None
    )

    reserved_names = {
        key
        for observation in observations
        if (key := _duplicate_key(observation.name)) is not None
    }

    resolved_observations: list[ExcelDeviceObservation] = []

    for observation in observations:
        row_number = observation.line_number

        if row_number is None:
            raise ValueError(
                "Excel observation is missing its source row number"
            )

        name_key = _duplicate_key(observation.name)
        serial_key = _serial_key(observation.serial_number)

        resolved = observation

        # Duplicate device names
        if name_key is not None and name_counts[name_key] > 1:
            if duplicate_name_policy == "error":
                raise ValueError(
                    f"Duplicate device name at row {row_number}: "
                    f"{observation.name}"
                )

            if duplicate_name_policy in {"prefix", "warn"}:
                issues.append(
                    InventoryIssue(
                        row_number=row_number,
                        reason=f"Duplicate device name: {observation.name}",
                    )
                )

            if duplicate_name_policy == "prefix":
                base_name = f"DUP-{row_number}-{observation.name}"
                generated_name = base_name
                suffix = 2

                while generated_name.strip().casefold() in reserved_names:
                    generated_name = f"{base_name}-{suffix}"
                    suffix += 1

                reserved_names.add(generated_name.strip().casefold())

                resolved = replace(
                    observation,
                    name=generated_name,
                )

        # Duplicate serial numbers
        if serial_key is not None and serial_counts[serial_key] > 1:
            if duplicate_serial_policy == "error":
                raise ValueError(
                    f"Duplicate serial number at row {row_number}: "
                    f"{observation.serial_number}"
                )

            if duplicate_serial_policy == "warn":
                issues.append(
                    InventoryIssue(
                        row_number=row_number,
                        reason=(
                            "Duplicate serial number: "
                            f"{observation.serial_number}"
                        ),
                    )
                )

        resolved_observations.append(resolved)

    return resolved_observations


# ---------------------------------------------------------------------------
# Excel row parsing
# ---------------------------------------------------------------------------

def _observation_from_row(
    row: tuple[Any, ...],
    *,
    row_number: int,
    header: tuple[Any, ...],
    columns: Mapping[str, int],
    issues: list[InventoryIssue],
) -> ExcelDeviceObservation:
    """Parse a single workbook row into an ExcelDeviceObservation."""
    name = _text(_cell(row, columns, NAME_COLUMN))

    if name is None:
        issues.append(
            InventoryIssue(
                row_number=row_number,
                reason="Missing device name",
            )
        )

    site = _text(_cell(row, columns, SITE_COLUMN))

    if site is None:
        issues.append(
            InventoryIssue(
                row_number=row_number,
                reason="Missing site",
            )
        )

    management_ip = _ipv4(
        _cell(row, columns, MANAGEMENT_IP_COLUMN),
        row_number=row_number,
        field_name="management IPv4 address",
        issues=issues,
        required=True,
    )

    primary_ipv4 = _ipv4(
        _cell(row, columns, PRIMARY_IPV4_COLUMN),
        row_number=row_number,
        field_name="primary IPv4 address",
        issues=issues,
    )

    return ExcelDeviceObservation(
        name=name,
        host_name=_text(_cell(row, columns, HOST_NAME_COLUMN)),
        site=site,
        role=_text(_cell(row, columns, ROLE_COLUMN)),
        primary_IPv4=primary_ipv4,
        management_ip=management_ip,
        manufacturer=_text(_cell(row, columns, MANUFACTURER_COLUMN)),
        model=_text(_cell(row, columns, MODEL_COLUMN)),
        serial_number=_text(_cell(row, columns, SERIAL_NUMBER_COLUMN)),
        asset_tag=_text(_cell(row, columns, ASSET_TAG_COLUMN)),
        software_version=_text(
            _cell(row, columns, SOFTWARE_VERSION_COLUMN)
        ),
        status=_text(_cell(row, columns, STATUS_COLUMN)),
        validation_status=_text(
            _cell(row, columns, VALIDATION_STATUS_COLUMN)
        ),
        line_number=row_number,
        raw_line=_raw_line(row, header),
    )


# ---------------------------------------------------------------------------
# Detailed Excel inventory reader
# ---------------------------------------------------------------------------

def read_excel_inventory(
    path: Path,
    sheet_name: str | None = None,
    *,
    duplicate_name_policy: DuplicateNamePolicy = "prefix",
    duplicate_serial_policy: DuplicateSerialPolicy = "warn",
) -> InventoryReadResult:
    """Read and validate Excel inventory without canonical reconciliation.

    This function retains the original detailed parsing contract.

    It returns:
        - ExcelDeviceObservation records
        - InventoryIssue records
        - Management IP-to-site mappings

    All non-empty rows are retained, including rows with invalid fields.

    Duplicate checks apply to the complete workbook, not a filtered subset.
    """
    if duplicate_name_policy not in {
        "prefix", "warn", "error", "ignore"
    }:
        raise ValueError(
            f"Invalid duplicate name policy: {duplicate_name_policy}"
        )

    if duplicate_serial_policy not in {
        "warn", "error", "ignore"
    }:
        raise ValueError(
            f"Invalid duplicate serial policy: {duplicate_serial_policy}"
        )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = (
            workbook[sheet_name]
            if sheet_name is not None
            else workbook.active
        )

        if sheet is None:
            raise ValueError("Missing inventory sheet")

        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)

        if header is None:
            raise ValueError("Inventory sheet is empty")

        columns = {
            text: index
            for index, value in enumerate(header)
            if (text := _text(value)) is not None
        }

        required = {
            NAME_COLUMN,
            SITE_COLUMN,
            MANAGEMENT_IP_COLUMN,
        }

        missing = required - columns.keys()

        if missing:
            raise ValueError(
                "Missing inventory columns: "
                f"{', '.join(sorted(missing))}"
            )

        observations: list[ExcelDeviceObservation] = []
        issues: list[InventoryIssue] = []
        sites_by_ip: dict[IPv4Address, set[str]] = {}

        for row_number, row in enumerate(rows, start=2):
            if all(_text(value) is None for value in row):
                continue

            observation = _observation_from_row(
                row,
                row_number=row_number,
                header=header,
                columns=columns,
                issues=issues,
            )

            observations.append(observation)

            if (
                observation.management_ip is None
                or observation.site is None
            ):
                continue

            known_sites = sites_by_ip.setdefault(
                observation.management_ip,
                set(),
            )

            if known_sites and observation.site not in known_sites:
                issues.append(
                    InventoryIssue(
                        row_number=row_number,
                        reason=(
                            "Management IP appears under multiple sites"
                        ),
                    )
                )

            known_sites.add(observation.site)

        observations = _resolve_duplicate_observations(
            observations,
            issues,
            duplicate_name_policy=duplicate_name_policy,
            duplicate_serial_policy=duplicate_serial_policy,
        )

        result = InventoryReadResult(
            observations=tuple(observations),
            issues=tuple(issues),
            ip_sites_mapping=MappingProxyType({
                ip: frozenset(sites)
                for ip, sites in sites_by_ip.items()
            }),
        )

        logger.info(
            "Excel inventory parsed",
            extra={
                "observation_count": len(result.observations),
                "issue_count": len(result.issues),
                "distinct_management_ips": len(sites_by_ip),
            },
        )

        return result

    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# SourceDataset normalization
# ---------------------------------------------------------------------------

def _to_device_observation(
    row: ExcelDeviceObservation,
    *,
    sheet_name: str | None,
) -> DeviceObservation:
    """Convert an Excel row into a normalized DeviceObservation.

    Source-reported attributes are preserved without asserting canonical
    identity or applying field-level reconciliation policy.
    """
    if row.line_number is None:
        raise ValueError(
            "Excel observation is missing its source row number"
        )

    sheet_ref = sheet_name or "<active>"

    return DeviceObservation(
        local_ref=f"excel:{sheet_ref}:row:{row.line_number}",
        inventory_name=row.name,
        hostname=row.host_name,
        serial_number=row.serial_number,
        asset_tag=row.asset_tag,
        management_ip=(
            str(row.management_ip)
            if row.management_ip is not None
            else None
        ),
        primary_ipv4=(
            str(row.primary_IPv4)
            if row.primary_IPv4 is not None
            else None
        ),
        manufacturer=row.manufacturer,
        model=row.model,
        software_version=row.software_version,
        role=row.role,
        site=row.site,
        status=row.status,
        validation_status=row.validation_status,
    )


def _selected(
    row: ExcelDeviceObservation,
    *,
    sites: frozenset[str] | None,
    devices: frozenset[str] | None,
) -> bool:
    """Check optional site and device filters.

    When both filters are supplied, both must match.
    Device selection supports inventory name, hostname and serial.
    """
    if sites is not None:
        site_key = _duplicate_key(row.site)

        if site_key not in sites:
            return False

    if devices is not None:
        identifiers = {
            key
            for value in (
                row.name,
                row.host_name,
                row.serial_number,
            )
            if (key := _duplicate_key(value)) is not None
        }

        if not identifiers.intersection(devices):
            return False

    return True


# ---------------------------------------------------------------------------
# Public normalized importer
# ---------------------------------------------------------------------------

def read_network_devices(
    path: Path,
    sheet_name: str | None = None,
    *,
    duplicate_name_policy: DuplicateNamePolicy = "prefix",
    duplicate_serial_policy: DuplicateSerialPolicy = "warn",
    sites: Collection[str] | None = None,
    devices: Collection[str] | None = None,
) -> SourceDataset:
    """Import Excel device inventory into a normalized SourceDataset.

    The entire workbook is parsed and validated before filters are applied.

    Args:
        path:
            Excel workbook path.

        sheet_name:
            Optional worksheet name. Defaults to the active worksheet.

        duplicate_name_policy:
            Handling of duplicated device inventory names.

        duplicate_serial_policy:
            Handling of duplicated non-placeholder serial numbers.

        sites:
            Optional site selection. Matching is case-insensitive.

        devices:
            Optional device selection by inventory name, hostname or serial.

    Returns:
        SourceDataset containing independently normalized device observations.

    Notes:
        - No canonical NetworkDevice objects are created.
        - No device identity decisions are made.
        - No SQL persistence is performed.
        - Filtering does not bypass workbook-wide duplicate validation.
        - Source identity is stable across repeated imports.
        - Each import receives independent observation and dataset UUIDs.
    """
    inventory = read_excel_inventory(
        path,
        sheet_name=sheet_name,
        duplicate_name_policy=duplicate_name_policy,
        duplicate_serial_policy=duplicate_serial_policy,
    )

    selected_sites = _normalize_filter(sites)
    selected_devices = _normalize_filter(devices)

    observations = tuple(
        _to_device_observation(
            row,
            sheet_name=sheet_name,
        )
        for row in inventory.observations
        if _selected(
            row,
            sites=selected_sites,
            devices=selected_devices,
        )
    )

    dataset = SourceDataset(
        source=Source(
            name=EXCEL_SOURCE_NAME,
            source_type=EXCEL_SOURCE_TYPE,
            source_id=EXCEL_SOURCE_ID,
        ),
        devices=observations,
    )

    logger.info(
        "Excel inventory normalized into SourceDataset",
        extra={
            "source_id": str(EXCEL_SOURCE_ID),
            "dataset_id": str(dataset.dataset_id),
            "total_inventory_rows": len(inventory.observations),
            "selected_observations": len(dataset.devices),
            "validation_issues": len(inventory.issues),
            "selected_site_count": (
                len(selected_sites)
                if selected_sites is not None
                else None
            ),
            "selected_device_count": (
                len(selected_devices)
                if selected_devices is not None
                else None
            ),
        },
    )

    return dataset
