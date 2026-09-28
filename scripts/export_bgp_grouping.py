"""Export grouped Palo Alto BGP prefixes to Excel and JSON.

Excel contains exactly two sheets:
- AS Numbers: vertically merged ASN cells, one prefix per row, blank Scope.
- Duplicated IPv4 Networks: merged prefix cells and every duplicate record.

Grouping is performed for one selected virtual-router section. Parsing
diagnostics and routes without a known origin are retained in JSON.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path



from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from netvoyager_network.bgp.grouping import group_as_prefixes
from netvoyager_network.bgp.models import AsGroupingResult, AsPrefix
from netvoyager_paloalto.bgp.models import BgpRibResult
from netvoyager_paloalto.bgp.parser import parse_loc_rib


NAVY = "17365D"
BLUE = "EAF1F8"
WHITE = "FFFFFF"
TEXT = "243746"
GRID = "CBD5E1"
INPUT = "FFF8E6"

THIN = Side(style="thin", color=GRID)
GROUP_EDGE = Side(style="thin", color=NAVY)


def prepare_sheet(
    sheet: Worksheet,
    headers: list[str],
    widths: tuple[float, ...],
    router_name: str,
) -> None:
    """Apply shared report formatting without adding extra data columns.

    Column headers remain in row one and repeat on printed pages.
    Sorting and filtering controls are omitted because the report uses
    vertically merged groups.
    """
    sheet.append(headers)
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 100
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    for cell, width in zip(sheet[1], widths):
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.column_dimensions[cell.column_letter].width = width

    sheet.row_dimensions[1].height = 28
    sheet.sheet_properties.tabColor = NAVY
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.outlinePr.summaryRight = False

    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.page_margins.top = 0.6
    sheet.page_margins.bottom = 0.6

    sheet.oddHeader.left.text = "NetVoyager | BGP Prefix Report"
    sheet.oddHeader.right.text = router_name.replace("&", "&&")
    sheet.oddFooter.left.text = sheet.title
    sheet.oddFooter.right.text = "Page &P of &N"


def format_group(
    sheet: Worksheet,
    first_row: int,
    last_row: int,
    group_index: int,
    *,
    blank_scope: bool = False,
) -> None:
    """Style a group and merge its first-column cells when needed.

    Only the group identifier is merged. Prefixes, ASNs in duplicate
    records, AS paths, and blank Scope cells remain individual cells.
    """
    background = BLUE if group_index % 2 == 0 else WHITE

    for row_number in range(first_row, last_row + 1):
        sheet.row_dimensions[row_number].height = 22

        for column in range(1, 4):
            cell = sheet.cell(row_number, column)
            cell.font = Font(name="Calibri", size=11, color=TEXT)
            cell.fill = PatternFill(
                "solid",
                fgColor=INPUT if blank_scope and column == 3 else background,
            )
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(
                left=THIN,
                right=THIN,
                top=GROUP_EDGE if row_number == first_row else THIN,
                bottom=GROUP_EDGE if row_number == last_row else THIN,
            )

    if last_row > first_row:
        sheet.merge_cells(
            start_row=first_row,
            start_column=1,
            end_row=last_row,
            end_column=1,
        )

    identifier = sheet.cell(first_row, 1)
    identifier.font = Font(
        name="Calibri", size=11, bold=True, color=NAVY
    )
    identifier.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


def export_excel(
    result: AsGroupingResult,
    destination: Path,
    router_name: str,
) -> None:
    """Write the agreed two-sheet report from an existing grouping result.

    ASNs and networks are sorted. Duplicate records preserve their input
    order and are never deduplicated again during export.

    Scope cells remain empty for later manual completion. Unknown paths
    display as 'Unknown'; an explicitly empty path displays as '(empty)'.
    """
    workbook = Workbook()
    workbook.properties.title = "BGP Prefix Grouping"
    workbook.properties.subject = f"Virtual router: {router_name}"
    workbook.properties.creator = "NetVoyager"

    systems_sheet = workbook.active
    systems_sheet.title = "AS Numbers"
    prepare_sheet(
        systems_sheet,
        ["ASN", "Prefix", "Scope"],
        (16, 28, 36),
        router_name,
    )

    for index, (asn, system) in enumerate(
        sorted(result.autonomous_systems.items())
    ):
        first_row = systems_sheet.max_row + 1

        for prefix in sorted(system.prefixes):
            systems_sheet.append([None, str(prefix), None])

        # Preserve an empty AS object if an exporter caller supplies one.
        if systems_sheet.max_row < first_row:
            systems_sheet.append([asn, None, None])

        systems_sheet.cell(first_row, 1, asn)
        format_group(
            systems_sheet,
            first_row,
            systems_sheet.max_row,
            index,
            blank_scope=True,
        )

    duplicates_sheet = workbook.create_sheet("Duplicated IPv4 Networks")
    prepare_sheet(
        duplicates_sheet,
        ["IPv4Network", "asn", "as_path"],
        (26, 16, 54),
        router_name,
    )

    for index, (prefix, records) in enumerate(
        sorted(result.duplicates.items())
    ):
        first_row = duplicates_sheet.max_row + 1

        for record in records:
            if record.as_path is None:
                path_text = "Unknown"
            elif not record.as_path:
                path_text = "(empty)"
            else:
                path_text = " ".join(map(str, record.as_path))

            duplicates_sheet.append([None, record.asn, path_text])

            # Increase height for long paths while retaining wrapped text.
            lines = max(1, (len(path_text) + 44) // 45)
            duplicates_sheet.row_dimensions[
                duplicates_sheet.max_row
            ].height = min(409, 18 * lines + 4)

        if not records:
            continue

        duplicates_sheet.cell(first_row, 1, str(prefix))

        # Preserve the path heights while applying common group formatting.
        heights = {
            row: duplicates_sheet.row_dimensions[row].height
            for row in range(first_row, duplicates_sheet.max_row + 1)
        }
        format_group(
            duplicates_sheet,
            first_row,
            duplicates_sheet.max_row,
            index,
        )
        for row, height in heights.items():
            duplicates_sheet.row_dimensions[row].height = height

    for sheet in workbook.worksheets:
        sheet.print_area = f"A1:C{sheet.max_row}"

    try:
        workbook.save(destination)
    finally:
        workbook.close()


def grouping_payload(result: AsGroupingResult) -> dict:
    """Convert network objects into explicit JSON-compatible values.

    ASN dictionary keys become strings, networks become CIDR strings,
    and AS paths become arrays. None and an empty path remain distinct.
    """
    return {
        "autonomous_systems": {
            str(asn): {
                "asn": asn,
                "prefixes": [
                    str(prefix) for prefix in sorted(system.prefixes)
                ],
            }
            for asn, system in sorted(result.autonomous_systems.items())
        },
        "duplicates": {
            str(prefix): [
                {
                    "asn": record.asn,
                    "prefix": str(record.prefix),
                    "as_path": (
                        list(record.as_path)
                        if record.as_path is not None
                        else None
                    ),
                }
                for record in records
            ]
            for prefix, records in sorted(result.duplicates.items())
        },
    }


def select_router(
    routers: tuple[BgpRibResult, ...],
    name: str | None,
) -> BgpRibResult:
    """Select one section without silently combining routing tables.

    A single section is selected automatically. Multiple sections require
    an explicit name that identifies exactly one section.
    """
    if name is None and len(routers) == 1:
        return routers[0]

    matches = [router for router in routers if router.virtual_router == name]

    if name is not None and len(matches) == 1:
        return matches[0]

    available = ", ".join(router.virtual_router for router in routers)
    raise ValueError(
        "Select exactly one section with --virtual-router. "
        f"Available sections: {available}"
    )


def main() -> None:
    """Read a RIB file, group its prefixes, and write Excel and JSON exports.

    The rightmost ASN in a populated AS path supplies the origin ASN.
    Routes without a path are retained in JSON but excluded from grouping.
    Rejected lines remain visible in JSON and console diagnostics.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Palo Alto RIB text file")
    parser.add_argument("--virtual-router", help="Virtual-router name")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/bgp-grouping"),
        help="Output path prefix without an extension",
    )
    args = parser.parse_args()

    dump = parse_loc_rib(args.input.read_text(encoding="utf-8-sig"))

    try:
        router = select_router(dump.routers, args.virtual_router)
    except ValueError as exc:
        parser.error(str(exc))

    records = [
        AsPrefix(
            asn=route.as_path[-1],
            prefix=route.prefix,
            as_path=route.as_path,
        )
        for route in router.routes
        if route.as_path
    ]

    unknown_origins = [
        {
            "prefix": str(route.prefix),
            "line_number": route.line_number,
            "raw_line": route.raw_line,
        }
        for route in router.routes
        if not route.as_path
    ]

    unparsed = [
        {
            "line_number": item.line_number,
            "reason": item.reason,
            "raw_line": item.raw_line,
        }
        for item in (*dump.unparsed_lines, *router.unparsed_lines)
    ]

    result = group_as_prefixes(records)

    report = {
        "schema_version": "1.0",
        "source_file": str(args.input),
        "virtual_router": router.virtual_router,
        "virtual_router_id": router.virtual_router_id,
        "parsed_routes": len(router.routes),
        "reported_routes": router.total_routes_shown,
        "grouped_records": len(records),
        **grouping_payload(result),
        "unknown_origins": unknown_origins,
        "unparsed_lines": unparsed,
    }

    excel_path = Path(f"{args.output}.xlsx")
    json_path = Path(f"{args.output}.json")
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    print(
    f"Parsed {len(router.routes)} routes; "
    f"grouped into {len(result.autonomous_systems)} ASNs.",
    flush=True,
    )
    print(f"Writing Excel: {excel_path.resolve()}", flush=True)

    export_excel(result, excel_path, router.virtual_router)

    print("Excel saved. Writing JSON...", flush=True)
    
    json_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )

    print(f"Virtual router: {router.virtual_router}")
    print(f"Parsed routes: {len(router.routes)}")
    print(f"Grouped records: {len(records)}")
    print(f"Autonomous systems: {len(result.autonomous_systems)}")
    print(f"Duplicated prefixes: {len(result.duplicates)}")
    print(f"Unknown origins: {len(unknown_origins)}")
    print(f"Unparsed lines: {len(unparsed)}")

    for item in unparsed:
        print(f"  Line {item['line_number']}: {item['reason']}")

    reported = router.total_routes_shown
    if reported is not None and reported != len(router.routes):
        print(f"WARNING: source footer reports {reported} routes.")

    print(f"Excel: {excel_path.resolve()}")
    print(f"JSON:  {json_path.resolve()}")


if __name__ == "__main__":
    main()